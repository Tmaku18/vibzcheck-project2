"""Generate the rubric-required "questions only" Word document.

The Project 2 rubric (Curated Questions tab) requires a separate Word
document submitted *before* the presentation that contains the 12-15
selected questions and **no answers**. This script generates that file
from the same 14 selections documented in `CURATED_QUESTIONS.md`,
keeping the two perfectly in sync.

Run:
    python submission/make_questions_only_doc.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_PATH = Path(__file__).resolve().parent / (
    "Makuvaza_Tanaka_Vibzcheck_Project2_Curated_Questions.docx"
)

HEADER = {
    "Author": "Tanaka Makuvaza",
    "Course": "CSC 6370 — Mobile App Development",
    "CRN": "13598",
    "Student ID": "002252191",
    "Project": "Vibzcheck — Final Project (Project 2), Solo (M.S. graduate track)",
}

# Each tuple is (category, question text). Mirrors the headings in
# CURATED_QUESTIONS.md exactly so the answer key and the question list
# never drift apart. 14 questions total: Implementation x3,
# Architecture x3, Testing x3, Firebase x4, Reflection x1.
QUESTIONS: list[tuple[str, str]] = [
    (
        "Implementation",
        "Feature Build Sequence: Describe the exact order you implemented "
        "your three most complex mobile features and why that order reduced risk.",
    ),
    (
        "Implementation",
        "Feature Build Sequence: Which feature was reworked after user-flow "
        "testing, and what changed in code and UI behavior?",
    ),
    (
        "Implementation",
        "State and Data Synchronization: Explain one real synchronization bug "
        "you encountered between app state and backend/local data. What was the "
        "root cause and what safeguards now prevent regression?",
    ),
    (
        "Architecture",
        "Navigation and Screen Responsibility: Choose one screen that became "
        "too large and explain how you decomposed responsibilities.",
    ),
    (
        "Architecture",
        "Navigation and Screen Responsibility: How does your current navigation "
        "structure support maintainability and future feature additions?",
    ),
    (
        "Architecture",
        "Security and Data Decisions: Identify one Firebase Security Rule "
        "(using request.auth.uid or custom claims) that directly prevented a "
        "bad write or unauthorized read - show the rule and the blocked "
        "scenario. Describe the trade-off between strict Firestore rule "
        "validation and iterative development speed in your project.",
    ),
    (
        "Testing",
        "Failure Case Ownership: Show one failure case your first "
        "implementation missed (network, auth, null state, or lifecycle "
        "issue). How did you redesign the UX response so users can recover "
        "without confusion?",
    ),
    (
        "Testing",
        "Failure Case Ownership: Show a second failure case (auth) you "
        "redesigned the UX around so users recover gracefully without leaking "
        "internal error codes.",
    ),
    (
        "Testing",
        "Performance Under Real Use: Which app interaction had noticeable lag, "
        "and how did you profile it? What code-level optimization gave the "
        "biggest user-visible improvement?",
    ),
    (
        "Firebase",
        "Firestore Data Modeling: Walk through your Firestore collection "
        "hierarchy. Why did you choose subcollections over top-level "
        "collections for your primary relational data?",
    ),
    (
        "Firebase",
        "Firestore Data Modeling: Which compound query required a composite "
        "index, and how does the index impact read cost as the collection grows?",
    ),
    (
        "Firebase",
        "Firebase Authentication & Security Rules: Show a Security Rule that "
        "uses request.auth.uid or custom claims to scope access. Explain what "
        "happens when an unauthenticated request hits that path.",
    ),
    (
        "Firebase",
        "Firebase Cloud Messaging (FCM): Describe the full FCM token lifecycle "
        "- when it is requested, where it is stored in Firestore, how it is "
        "refreshed, and how a notification is delivered. How does your app "
        "handle notification delivery differently for foreground vs. "
        "background vs. terminated app states?",
    ),
    (
        "Reflection",
        "Team Engineering Reflection: What decision from early planning created "
        "technical debt later, and how did you resolve it? If you restarted "
        "this app tomorrow, what would you architect differently first?",
    ),
]


def build_document() -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Vibzcheck - Curated Questions (Selection)", level=0)
    title.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER

    for label, value in HEADER.items():
        para = doc.add_paragraph()
        run = para.add_run(f"{label}: ")
        run.bold = True
        para.add_run(value)

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.add_run(
        "The 14 questions below are my pre-presentation selections from the "
        "official Curated Questions tab. They are submitted with no answers "
        "per the rubric. Coverage: Implementation x3, Architecture x3, "
        "Testing x3, Firebase x4, Reflection x1."
    )
    doc.add_paragraph()

    counts: dict[str, int] = {}
    current_section: str | None = None
    overall = 0

    for category, question in QUESTIONS:
        if category != current_section:
            doc.add_heading(category, level=1)
            current_section = category
        counts[category] = counts.get(category, 0) + 1
        overall += 1
        para = doc.add_paragraph(style="List Number")
        para.add_run(f"({overall}) ").bold = True
        para.add_run(question)

    doc.add_paragraph()
    summary = doc.add_paragraph()
    summary.add_run("Selected total: ").bold = True
    summary.add_run(
        ", ".join(f"{k} x{v}" for k, v in counts.items())
        + f" - {sum(counts.values())} questions."
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
